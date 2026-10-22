from pathlib import Path
from docx import Document

source = Path(r"D:\học tập\tài liệu\redeploy code lên azure.docx")
output = Path(r"D:\h-solution-project\redeploy code lên azure h-solution-project.docx")

doc = Document(str(source))
paragraphs = doc.paragraphs

content = [
    ("Azure Deployment Guide - H-Solution Backend FastAPI Docker", "Title"),
    ("1. Build and Push Docker Image to ACR", "Heading 1"),
    ("az acr build --registry multitaskingcommunity --image h-solution-api:latest ./backend", "Normal"),
    ("2. Configure Container App Registry Access", "Heading 1"),
    ("az containerapp registry set --name h-solution-project --resource-group Multitasking-Community --server multitaskingcommunity.azurecr.io --identity system", "Normal"),
    ("3. Update Container App Image", "Heading 1"),
    ("az containerapp update --name h-solution-project --resource-group Multitasking-Community --image multitaskingcommunity.azurecr.io/h-solution-api:latest", "Normal"),
    ("4. Verify Deployment", "Heading 1"),
    ("Invoke-WebRequest -Uri https://h-solution-project.greenplant-669971f3.southeastasia.azurecontainerapps.io/health", "Normal"),
    ("5. Restart Revision (optional)", "Heading 1"),
    ("az containerapp revision restart --name h-solution-project --resource-group Multitasking-Community", "Normal"),
    ("6. Current Container App Configuration", "Heading 1"),
    ("CPU: 0.25 vCPU | Memory: 0.5 GiB | Min replicas: 1 | Max replicas: 1 | Target port: 8000", "Normal"),
    ("", "Normal"),
    ("", "Normal"),
    ("Azure Deployment Guide - H-Solution Backend Configuration", "Title"),
    ("1. Container App URL", "Heading 1"),
    ("https://h-solution-project.greenplant-669971f3.southeastasia.azurecontainerapps.io", "Normal"),
    ("2. Docker Image", "Heading 1"),
    ("multitaskingcommunity.azurecr.io/h-solution-api:latest", "Normal"),
    ("3. Health Check", "Heading 1"),
    ("GET /health returns {\"status\": \"ok\"} when the backend is ready.", "Normal"),
    ("4. Required Environment Variables", "Heading 1"),
    ("SUPABASE_URL, SUPABASE_KEY, JWT_SECRET_KEY, and the configured payment, map, Firebase, and OTP variables must be set in the Container App environment. Do not put secret values in this document or in the Docker image.", "Normal"),
    ("5. Redeploy Workflow", "Heading 1"),
    ("Update the backend source, run the ACR build command, update the Container App image, and verify /health before testing API routes.", "Normal"),
]

def replace_paragraph(paragraph, text):
    for run in paragraph.runs:
        run.text = ""
    if paragraph.runs:
        paragraph.runs[0].text = text
    else:
        paragraph.add_run(text)

for paragraph, (text, style) in zip(paragraphs, content):
    replace_paragraph(paragraph, text)
    paragraph.style = style

doc.save(str(output))
print(output)
